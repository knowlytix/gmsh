"""1-page slide: Testing Execution.

Grounded in gmsh_testing_ext chapters 11-13 (SUT Adapters & Traces,
Tool Gateway & Fault Injection, Policy Engine Runtime Enforcement).

Three-panel arc, using monograph's actual terminology:
  SUTTrace / AgentAdapter   ->   ToolGateway / FaultProfile   ->   PolicyEngine 4-layer pipeline.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Colors --------------------------------------------------------------
NAVY   = RGBColor(0x14, 0x2B, 0x4D)
DARK   = RGBColor(0x22, 0x26, 0x2E)
BODY   = RGBColor(0x3E, 0x44, 0x50)
MID    = RGBColor(0x6C, 0x72, 0x80)
FAINT  = RGBColor(0xD3, 0xD7, 0xDE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

RED    = RGBColor(0xB2, 0x3A, 0x2F)   # SUTTrace (observe)
BLUE   = RGBColor(0x1F, 0x6F, 0xB5)   # ToolGateway (intervene)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)   # PolicyEngine (enforce)

RED_BG    = RGBColor(0xFB, 0xE9, 0xE7)
BLUE_BG   = RGBColor(0xE6, 0xF0, 0xFA)
GREEN_BG  = RGBColor(0xE7, 0xF4, 0xEC)

MONO = "Consolas"
SANS = "Calibri"


def txt(slide, left, top, width, height, runs,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        margin=(0.02, 0.02, 0.02, 0.02)):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left   = Inches(margin[0])
    tf.margin_right  = Inches(margin[1])
    tf.margin_top    = Inches(margin[2])
    tf.margin_bottom = Inches(margin[3])
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, italic, color, font in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = font
    return tb


def box(slide, left, top, width, height, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        left, top, width, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def arrow(slide, left, top, width, height, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


# --- Slide ---------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Navy stripe
box(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.10), NAVY)

# Title
txt(slide, Inches(0.5), Inches(0.22), Inches(12.4), Inches(0.55), [[
    ("Testing Execution", 30, True, False, NAVY, SANS),
    ("   —   traces, fault injection, policy enforcement",
     17, False, False, MID, SANS),
]])
txt(slide, Inches(0.5), Inches(0.78), Inches(12.4), Inches(0.28), [[
    ("Essay 6  ·  Agus Sudjianto  &  Aldo Lau", 11, False, True, MID, SANS),
]])

# One-line framing
txt(slide, Inches(0.5), Inches(1.18), Inches(12.4), Inches(0.42), [[
    ("Reasoning tests examine what the agent ",
     13.5, False, False, BODY, SANS),
    ("said",
     14, True, True, DARK, SANS),
    (".  Execution tests examine what the agent ",
     13.5, False, False, BODY, SANS),
    ("did",
     14, True, True, BLUE, SANS),
    (".  Three layers, installed at harness hooks, non-invasive to the SUT.",
     13.5, False, False, BODY, SANS),
]])

# -- Three-panel concept arc ----------------------------------------------
t_top   = Inches(1.78)
t_h     = Inches(3.75)
panel_w = Inches(3.85)
gap     = Inches(0.45)
start_l = Inches(0.5 + (12.4 - (3 * 3.85 + 2 * 0.45)) / 2.0)


def make_panel(left, accent, accent_bg, title, tagline, body_runs, footer_runs):
    box(slide, left, t_top, panel_w, t_h, fill=accent_bg, line=accent, rounded=True)
    box(slide, left, t_top, panel_w, Inches(0.45), fill=accent, rounded=True)
    box(slide, left, t_top + Inches(0.33), panel_w, Inches(0.12), fill=accent)
    txt(slide, left + Inches(0.15), t_top + Inches(0.04),
        panel_w - Inches(0.3), Inches(0.38), [[
            (title, 16, True, False, WHITE, SANS),
            ("   " + tagline, 11, False, True, WHITE, SANS),
        ]], anchor=MSO_ANCHOR.MIDDLE)
    box(slide, left + Inches(0.25), t_top + Inches(0.65),
        panel_w - Inches(0.5), Inches(2.55), fill=WHITE, line=FAINT, rounded=True)
    txt(slide, left + Inches(0.38), t_top + Inches(0.78),
        panel_w - Inches(0.76), Inches(2.42), body_runs)
    txt(slide, left + Inches(0.20), t_top + t_h - Inches(0.42),
        panel_w - Inches(0.4), Inches(0.35), footer_runs,
        align=PP_ALIGN.CENTER)


# ---------- Panel 1: SUTTrace / Adapter ----------------------------------
l1 = start_l
make_panel(
    l1, RED, RED_BG,
    "SUT trace",
    "adapter captures what ran",
    [
        [("Wrap the agent in an ",
          11.5, False, False, BODY, SANS),
         ("adapter",
          11.5, True, False, DARK, SANS),
         (".  Record every run as an ",
          11.5, False, False, BODY, SANS),
         ("SUTTrace",
          11.5, True, False, DARK, MONO),
         (".",
          11.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Each trace carries:",
          11, True, False, RED, SANS)],
        [("  ·  ",           10.5, False, False, DARK, SANS),
         ("tool_calls",      10.5, True,  False, DARK, MONO),
         ("   name, args, result, per-call latency",
          10.5, False, True, BODY, SANS)],
        [("  ·  ",           10.5, False, False, DARK, SANS),
         ("latency_ms",      10.5, True,  False, DARK, MONO),
         (" · ",             10.5, False, False, DARK, SANS),
         ("token_count",     10.5, True,  False, DARK, MONO)],
        [("  ·  ",           10.5, False, False, DARK, SANS),
         ("retrieved_context",10.5, True,  False, DARK, MONO),
         (" · ",             10.5, False, False, DARK, SANS),
         ("error",           10.5, True,  False, DARK, MONO)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Installed at ",
          10.5, False, False, BODY, SANS),
         ("harness hooks",
          10.5, True, False, DARK, SANS),
         (" (",
          10.5, False, False, BODY, SANS),
         ("on_tool_call",
          10.5, False, False, DARK, MONO),
         (", ",
          10.5, False, False, BODY, SANS),
         ("on_tool_result",
          10.5, False, False, DARK, MONO),
         (")  —  re-entrant, non-invasive.",
          10.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Every claim about what the agent ",
          10.5, False, False, BODY, SANS),
         ("did",
          10.5, True, True, RED, SANS),
         (" is grounded in the trace.",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("one uniform interface across LLM · agent · RAG",
         10.5, True, True, RED, SANS),
    ]],
)

# arrow 1
arrow(slide, l1 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), BLUE)

# ---------- Panel 2: ToolGateway / FaultProfile --------------------------
l2 = l1 + panel_w + gap
make_panel(
    l2, BLUE, BLUE_BG,
    "Tool gateway",
    "intercept · inject · observe",
    [
        [("A ",
          11.5, False, False, BODY, SANS),
         ("ToolGateway",
          11.5, True, False, DARK, MONO),
         (" sits between agent and tool and applies a ",
          11.5, False, False, BODY, SANS),
         ("FaultProfile",
          11.5, True, False, DARK, MONO),
         (".",
          11.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Per-call decision:",
          11, True, False, BLUE, SANS)],
        [("  ",        10.5, False, False, DARK, SANS),
         ("allow",     10.5, True,  False, GREEN, MONO),
         (" · ",       10.5, False, False, DARK, SANS),
         ("block",     10.5, True,  False, RED, MONO),
         (" · ",       10.5, False, False, DARK, SANS),
         ("inject_fault",10.5, True, False, BLUE, MONO)],
        [("  ",        10.5, False, False, DARK, SANS),
         ("mock_response",10.5, True, False, BLUE, MONO),
         (" · ",       10.5, False, False, DARK, SANS),
         ("stale_data",10.5, True,  False, BLUE, MONO)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Profile fields:  ",
          10.5, False, False, BODY, SANS),
         ("tool_pattern",
          10.5, False, False, DARK, MONO),
         (" · ",
          10.5, False, False, BODY, SANS),
         ("error_rate",
          10.5, False, False, DARK, MONO),
         (" · ",
          10.5, False, False, BODY, SANS),
         ("latency_ms",
          10.5, False, False, DARK, MONO),
         (".",
          10.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Observe recovery:  ",
          11, True, False, DARK, SANS),
         ("retry",
          11, True, False, BLUE, MONO),
         (" · ",
          11, False, False, BODY, SANS),
         ("fallback",
          11, True, False, BLUE, MONO),
         (" · ",
          11, False, False, BODY, SANS),
         ("escalate",
          11, True, False, BLUE, MONO),
         (".",
          11, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Every decision logged to a ",
          10.5, False, False, BODY, SANS),
         ("GatewayTranscript",
          10.5, True, False, DARK, MONO),
         (".",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("seeded, replayable, reconciles with the trace",
         10.5, True, True, BLUE, SANS),
    ]],
)

# arrow 2
arrow(slide, l2 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), GREEN)

# ---------- Panel 3: PolicyEngine four-layer pipeline --------------------
l3 = l2 + panel_w + gap
make_panel(
    l3, GREEN, GREEN_BG,
    "Policy engine",
    "four-layer pre-execution check",
    [
        [("Gateway asks ",
          11, False, False, BODY, SANS),
         ("can this succeed",
          11, True, True, DARK, SANS),
         (".  Policy asks ",
          11, False, False, BODY, SANS),
         ("should it be permitted",
          11, True, True, GREEN, SANS),
         (".",
          11, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Four layers, in order:",
          11, True, False, GREEN, SANS)],
        [("  1.  ",           10.5, False, False, DARK, SANS),
         ("Explicit rules",   10.5, True,  False, DARK, SANS),
         ("    fnmatch on ",  10.5, False, True, BODY, SANS),
         ("PolicyRule",       10.5, False, False, DARK, MONO),
         ("s",                10.5, False, False, BODY, SANS)],
        [("  2.  ",           10.5, False, False, DARK, SANS),
         ("Graph lookup",     10.5, True,  False, DARK, SANS),
         ("      implicit rules in GMS",
          10.5, False, True, BODY, SANS)],
        [("  3.  ",           10.5, False, False, DARK, SANS),
         ("Geometric",        10.5, True,  False, DARK, SANS),
         ("           calibrated plausibility",
          10.5, False, True, BODY, SANS)],
        [("  4.  ",           10.5, False, False, DARK, SANS),
         ("Contradiction",    10.5, True,  False, DARK, SANS),
         ("      tension energy on rule",
          10.5, False, True, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Rule types:  ",
          10.5, False, False, BODY, SANS),
         ("requires",
          10.5, True, False, DARK, MONO),
         (" · ",
          10.5, False, False, BODY, SANS),
         ("blocks",
          10.5, True, False, DARK, MONO),
         (" · ",
          10.5, False, False, BODY, SANS),
         ("allows",
          10.5, True, False, DARK, MONO),
         (".",
          10.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Same GMS primitives that verify facts now ",
          10.5, False, False, BODY, SANS),
         ("enforce governance",
          10.5, True, True, GREEN, SANS),
         (".",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("one policy store  ·  one audit log  ·  one calibration",
         10.5, True, True, GREEN, SANS),
    ]],
)

# -- Concrete result strip (520-question E2E) -----------------------------
dd_top = Inches(5.70)
box(slide, Inches(0.5), dd_top, Inches(12.4), Inches(0.55), WHITE,
    line=FAINT, rounded=True)
txt(slide, Inches(0.7), dd_top + Inches(0.07), Inches(3.1), Inches(0.40), [[
    ("520-question E2E run:", 11, True, False, NAVY, SANS),
]], anchor=MSO_ANCHOR.MIDDLE)
txt(slide, Inches(3.7), dd_top + Inches(0.07), Inches(9.1), Inches(0.40), [[
    ("1 733 gateway decisions   →   ",
     11, False, False, DARK, MONO),
    ("1 395 allow",
     11, True, False, GREEN, MONO),
    ("  ·  ",
     11, False, False, DARK, MONO),
    ("169 inject_fault",
     11, True, False, BLUE, MONO),
    ("  ·  ",
     11, False, False, DARK, MONO),
    ("169 policy block",
     11, True, False, RED, MONO),
    ("     ",
     11, False, False, DARK, MONO),
    ("(trace & transcript reconcile)",
     10, False, True, MID, MONO),
]], anchor=MSO_ANCHOR.MIDDLE)

# -- Bottom punchline bar -------------------------------------------------
bar_top = Inches(6.42)
box(slide, Inches(0.5), bar_top, Inches(12.4), Inches(0.95), NAVY, rounded=True)
txt(slide, Inches(0.75), bar_top + Inches(0.10), Inches(11.9), Inches(0.35), [[
    ("Observe with the adapter.  Stress with the gateway.  Enforce with the policy engine.  ",
     12, True, False, WHITE, SANS),
    ("Three layers, composable, all written to audit.",
     12, False, False, WHITE, SANS),
]])
txt(slide, Inches(0.75), bar_top + Inches(0.48), Inches(11.9), Inches(0.42), [[
    ("Every trace reconciles with every transcript.  Every policy block names the rule that fired.  ",
     11, False, True, WHITE, SANS),
    ("An agent that finishes on time is not an agent that acted correctly.",
     11, False, True, WHITE, SANS),
]])

out = "/home/asudjianto/jupyterlab/kg-memory/execution_slide.pptx"
prs.save(out)
print(f"Wrote: {out}")
