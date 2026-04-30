"""1-page slide: From Verdicts to Release Gates.

Concept first: a verdict is not a release decision.
Three-panel arc: the gap  ->  severity x risk tier  ->  the gate.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Colors --------------------------------------------------------------
NAVY   = RGBColor(0x14, 0x2B, 0x4D)
DARK   = RGBColor(0x22, 0x26, 0x2E)
BODY   = RGBColor(0x3E, 0x44, 0x50)
MID    = RGBColor(0x6C, 0x72, 0x80)
FAINT  = RGBColor(0xD3, 0xD7, 0xDE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

RED    = RGBColor(0xB2, 0x3A, 0x2F)   # the gap (problem)
BLUE   = RGBColor(0x1F, 0x6F, 0xB5)   # severity x risk (mechanism)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)   # the gate (decision)

RED_BG    = RGBColor(0xFB, 0xE9, 0xE7)
BLUE_BG   = RGBColor(0xE6, 0xF0, 0xFA)
GREEN_BG  = RGBColor(0xE7, 0xF4, 0xEC)

MONO = "Consolas"
SANS = "Calibri"


def txt(slide, left, top, width, height, runs,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        margin=(0.02, 0.02, 0.02, 0.02)):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left   = Inches(margin[0])
    tf.margin_right  = Inches(margin[1])
    tf.margin_top    = Inches(margin[2])
    tf.margin_bottom = Inches(margin[3])
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, italic, color, font in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = font
    return tb


def box(slide, left, top, width, height, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        left, top, width, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def arrow(slide, left, top, width, height, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


# --- Slide ---------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Navy stripe
box(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.10), NAVY)

# Title
txt(slide, Inches(0.5), Inches(0.22), Inches(12.4), Inches(0.55), [[
    ("From Verdicts to Release Gates", 30, True, False, NAVY, SANS),
    ("   —   severity, risk tiers, and provenance",
     17, False, False, MID, SANS),
]])
txt(slide, Inches(0.5), Inches(0.78), Inches(12.4), Inches(0.28), [[
    ("Essay 5  ·  Agus Sudjianto  &  Aldo Lau", 11, False, True, MID, SANS),
]])

# One-line framing
txt(slide, Inches(0.5), Inches(1.18), Inches(12.4), Inches(0.42), [[
    ("A verdict says this answer passed.  ",
     14, True, False, DARK, SANS),
    ("A release decision says this agent is ready to deploy in this role, under these risks.  ",
     13.5, False, False, BODY, SANS),
    ("Which failures matter, and by how much?",
     13.5, True, True, BLUE, SANS),
]])

# -- Three-panel concept arc ----------------------------------------------
t_top   = Inches(1.78)
t_h     = Inches(3.75)
panel_w = Inches(3.85)
gap     = Inches(0.45)
start_l = Inches(0.5 + (12.4 - (3 * 3.85 + 2 * 0.45)) / 2.0)


def make_panel(left, accent, accent_bg, title, tagline, body_runs, footer_runs):
    box(slide, left, t_top, panel_w, t_h, fill=accent_bg, line=accent, rounded=True)
    box(slide, left, t_top, panel_w, Inches(0.45), fill=accent, rounded=True)
    box(slide, left, t_top + Inches(0.33), panel_w, Inches(0.12), fill=accent)
    txt(slide, left + Inches(0.15), t_top + Inches(0.04),
        panel_w - Inches(0.3), Inches(0.38), [[
            (title, 16, True, False, WHITE, SANS),
            ("   " + tagline, 11, False, True, WHITE, SANS),
        ]], anchor=MSO_ANCHOR.MIDDLE)
    box(slide, left + Inches(0.25), t_top + Inches(0.65),
        panel_w - Inches(0.5), Inches(2.55), fill=WHITE, line=FAINT, rounded=True)
    txt(slide, left + Inches(0.38), t_top + Inches(0.78),
        panel_w - Inches(0.76), Inches(2.42), body_runs)
    txt(slide, left + Inches(0.20), t_top + t_h - Inches(0.42),
        panel_w - Inches(0.4), Inches(0.35), footer_runs,
        align=PP_ALIGN.CENTER)


# ---------- Panel 1: The gap (problem) -----------------------------------
l1 = start_l
make_panel(
    l1, RED, RED_BG,
    "The aggregation gap",
    "87% does not mean ship",
    [
        [("An agent scores ",
          11.5, False, False, BODY, SANS),
         ("87% on 520 tests",
          11.5, True, False, DARK, SANS),
         (".  Team ships.",
          11.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Three weeks later, a hallucination triggers an automated trade.",
          11, False, True, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("The failure ",
          11, False, False, BODY, SANS),
         ("was in the 13%",
          11, True, False, RED, SANS),
         (".  The aggregate number hid it.",
          11, False, False, BODY, SANS)],
        [(" ", 6, False, False, DARK, SANS)],
        [("Why it fails:",
          11.5, True, False, DARK, SANS)],
        [("  ·  treats every failure as fungible",
          10.5, False, False, BODY, SANS)],
        [("  ·  weights a hallucination the same as a paraphrase miss",
          10.5, False, False, BODY, SANS)],
        [("  ·  gives the release committee one number, not a decision",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("a score is not a ship / no-ship",
         11, True, True, RED, SANS),
    ]],
)

# arrow 1
arrow(slide, l1 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), BLUE)

# ---------- Panel 2: Severity x Risk Tier (mechanism) --------------------
l2 = l1 + panel_w + gap
make_panel(
    l2, BLUE, BLUE_BG,
    "Severity × risk tier",
    "same evidence, different stakes",
    [
        [("Each failure is classified by ",
          11, False, False, BODY, SANS),
         ("three inputs",
          11, True, False, DARK, SANS),
         (":",
          11, False, False, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("what broke",        10.5, True,  False, DARK, SANS),
         ("   (28 codes, 7 families)",
          10.5, False, True, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("where deployed",    10.5, True,  False, DARK, SANS),
         ("   (advisory / recommend / act)",
          10.5, False, True, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("how severe this one",10.5, True, False, DARK, SANS),
         ("   (geometric scalar)",
          10.5, False, True, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Same failure, escalated by tier:",
          11, True, False, BLUE, SANS)],
        [("   ",               10.5, False, False, DARK, MONO),
         ("fact_incorrect",    10.5, True,  False, DARK, MONO),
         ("      advisory  ",  10.5, False, False, BODY, MONO),
         ("HIGH",              10.5, True,  False, BODY, MONO)],
        [("                         action   ",
          10.5, False, False, BODY, MONO),
         ("CRITICAL",
          10.5, True,  False, RED, MONO)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Cost of being wrong differs by role.  The pipeline ",
          10.5, False, False, BODY, SANS),
         ("expresses that",
          10.5, True, False, DARK, SANS),
         (".",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("a research aide ≠ a loan approver",
         11, True, True, BLUE, SANS),
    ]],
)

# arrow 2
arrow(slide, l2 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), GREEN)

# ---------- Panel 3: The gate (decision) ---------------------------------
l3 = l2 + panel_w + gap
make_panel(
    l3, GREEN, GREEN_BG,
    "The release gate",
    "four dimensions, zero budgets",
    [
        [("Not one weighted sum  —  ",
          11, False, False, BODY, SANS),
         ("four independent gates",
          11, True, False, GREEN, SANS),
         (":",
          11, False, False, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("accuracy",          10.5, True,  False, DARK, SANS),
         ("            (how often correct)",
          10.5, False, True, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("geometric confidence",10.5, True, False, DARK, SANS),
         ("  (how aligned)",
          10.5, False, True, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("tension / holonomy", 10.5, True, False, DARK, SANS),
         ("     (consistency)",
          10.5, False, True, BODY, SANS)],
        [("   ·  ",            10.5, False, False, DARK, SANS),
         ("failure budget",    10.5, True,  False, DARK, SANS),
         ("          (zero-tolerance caps)",
          10.5, False, True, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Blocks are ",
          11, False, False, BODY, SANS),
         ("named, not scored",
          11, True, False, GREEN, SANS),
         (".  Every verdict carries:",
          11, False, False, BODY, SANS)],
        [("  ·  severity level  +  composite score",
          10.5, False, False, BODY, SANS)],
        [("  ·  audit packet  (Q · answer · factors · verdict)",
          10.5, False, False, BODY, SANS)],
        [("  ·  ",            10.5, False, False, BODY, SANS),
         ("provenance chain",  10.5, True,  False, DARK, SANS),
         ("  back to the source sentence",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("ship / no-ship with named reasons",
         11, True, True, GREEN, SANS),
    ]],
)

# -- Blocking-reason example (small strip) --------------------------------
dd_top = Inches(5.70)
box(slide, Inches(0.5), dd_top, Inches(12.4), Inches(0.55), WHITE,
    line=FAINT, rounded=True)
txt(slide, Inches(0.7), dd_top + Inches(0.07), Inches(2.7), Inches(0.40), [[
    ("Example block reason:", 11, True, False, NAVY, SANS),
]], anchor=MSO_ANCHOR.MIDDLE)
txt(slide, Inches(3.2), dd_top + Inches(0.07), Inches(9.5), Inches(0.40), [[
    ("[recommendation]  accuracy 0.73 < 0.80   ·   geo_conf 0.48 < 0.55   →   ",
     11, False, False, DARK, MONO),
    ("BLOCK",
     11, True, False, RED, MONO),
    ("   (traces to doc 97bf497d… §model_inventory 1432–1461)",
     10, False, True, MID, MONO),
]], anchor=MSO_ANCHOR.MIDDLE)

# -- Bottom punchline bar -------------------------------------------------
bar_top = Inches(6.42)
box(slide, Inches(0.5), bar_top, Inches(12.4), Inches(0.95), NAVY, rounded=True)
txt(slide, Inches(0.75), bar_top + Inches(0.10), Inches(11.9), Inches(0.35), [[
    ("Every threshold is calibrated.  Every blocking reason is named.  ",
     12, True, False, WHITE, SANS),
    ("Every verdict traces back to a specific offset in a specific document version.",
     12, False, False, WHITE, SANS),
]])
txt(slide, Inches(0.75), bar_top + Inches(0.48), Inches(11.9), Inches(0.42), [[
    ("The release committee does not need a system that is right most of the time  —  ",
     11, False, True, WHITE, SANS),
    ("it needs a system whose reasoning is inspectable, where every decision can be argued, overridden, or accepted on its merits.",
     11, False, True, WHITE, SANS),
]])

out = "/home/asudjianto/jupyterlab/kg-memory/release_gate_slide.pptx"
prs.save(out)
print(f"Wrote: {out}")
