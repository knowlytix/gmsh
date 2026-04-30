"""1-page slide: Judging Agent Output — Geometric Verification Without LLM Bias.

Concept first: score each question with the same primitive that generated it.
Three-panel arc: Trilemma fails  ->  Geometric Judge  ->  Head-to-head numbers.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Colors --------------------------------------------------------------
NAVY   = RGBColor(0x14, 0x2B, 0x4D)
DARK   = RGBColor(0x22, 0x26, 0x2E)
BODY   = RGBColor(0x3E, 0x44, 0x50)
MID    = RGBColor(0x6C, 0x72, 0x80)
FAINT  = RGBColor(0xD3, 0xD7, 0xDE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

RED    = RGBColor(0xB2, 0x3A, 0x2F)   # trilemma (problem)
BLUE   = RGBColor(0x1F, 0x6F, 0xB5)   # the judge (solution)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)   # head-to-head (result)

RED_BG    = RGBColor(0xFB, 0xE9, 0xE7)
BLUE_BG   = RGBColor(0xE6, 0xF0, 0xFA)
GREEN_BG  = RGBColor(0xE7, 0xF4, 0xEC)

MONO = "Consolas"
SANS = "Calibri"


def txt(slide, left, top, width, height, runs,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        margin=(0.02, 0.02, 0.02, 0.02)):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left   = Inches(margin[0])
    tf.margin_right  = Inches(margin[1])
    tf.margin_top    = Inches(margin[2])
    tf.margin_bottom = Inches(margin[3])
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, italic, color, font in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = font
    return tb


def box(slide, left, top, width, height, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        left, top, width, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def arrow(slide, left, top, width, height, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def pill(slide, left, top, width, height, fill, text_color, label, font_size=9.5):
    box(slide, left, top, width, height, fill, rounded=True)
    txt(slide, left, top, width, height, [[
            (label, font_size, True, False, text_color, SANS),
        ]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# --- Slide ---------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Navy stripe
box(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.10), NAVY)

# Title
txt(slide, Inches(0.5), Inches(0.22), Inches(12.4), Inches(0.55), [[
    ("Judging Agent Output", 30, True, False, NAVY, SANS),
    ("   —   geometric verification without LLM bias",
     17, False, False, MID, SANS),
]])
txt(slide, Inches(0.5), Inches(0.78), Inches(12.4), Inches(0.28), [[
    ("Essay 4  ·  Agus Sudjianto  &  Aldo Lau", 11, False, True, MID, SANS),
]])

# One-line framing
txt(slide, Inches(0.5), Inches(1.18), Inches(12.4), Inches(0.42), [[
    ("520 well-designed test cases are only useful if we can score them right.  ",
     14, True, False, DARK, SANS),
    ("The obvious answers fail in different ways.  We need a judge that is ",
     13.5, False, False, BODY, SANS),
    ("calibrated, deterministic, and auditable", 13.5, True, True, BLUE, SANS),
    (".", 13.5, False, False, BODY, SANS),
]])

# -- Three-panel concept arc ----------------------------------------------
t_top   = Inches(1.78)
t_h     = Inches(3.75)
panel_w = Inches(3.85)
gap     = Inches(0.45)
start_l = Inches(0.5 + (12.4 - (3 * 3.85 + 2 * 0.45)) / 2.0)

def make_panel(left, accent, accent_bg, title, tagline, body_runs, footer_runs):
    box(slide, left, t_top, panel_w, t_h, fill=accent_bg, line=accent, rounded=True)
    box(slide, left, t_top, panel_w, Inches(0.45), fill=accent, rounded=True)
    box(slide, left, t_top + Inches(0.33), panel_w, Inches(0.12), fill=accent)
    txt(slide, left + Inches(0.15), t_top + Inches(0.04),
        panel_w - Inches(0.3), Inches(0.38), [[
            (title, 16, True, False, WHITE, SANS),
            ("   " + tagline, 11, False, True, WHITE, SANS),
        ]], anchor=MSO_ANCHOR.MIDDLE)
    box(slide, left + Inches(0.25), t_top + Inches(0.65),
        panel_w - Inches(0.5), Inches(2.55), fill=WHITE, line=FAINT, rounded=True)
    txt(slide, left + Inches(0.38), t_top + Inches(0.78),
        panel_w - Inches(0.76), Inches(2.42), body_runs)
    txt(slide, left + Inches(0.20), t_top + t_h - Inches(0.42),
        panel_w - Inches(0.4), Inches(0.35), footer_runs,
        align=PP_ALIGN.CENTER)

# ---------- Panel 1: The trilemma (problem) ------------------------------
l1 = start_l
make_panel(
    l1, RED, RED_BG,
    "The trilemma",
    "three scorings, three failures",
    [
        [("Exact string match",
          12, True, False, DARK, SANS)],
        [("   ",   10, False, False, DARK, SANS),
         ("23.8%  —  brittle to paraphrase, reorder, hedge.",
          10.5, False, True, BODY, SANS)],
        [(" ", 6, False, False, DARK, SANS)],
        [("Knowledge-graph lookup",
          12, True, False, DARK, SANS)],
        [("   ",   10, False, False, DARK, SANS),
         ("Breaks on relational and multi-hop questions.",
          10.5, False, True, BODY, SANS)],
        [(" ", 6, False, False, DARK, SANS)],
        [("LLM-as-judge",
          12, True, False, DARK, SANS)],
        [("   ",   10, False, False, DARK, SANS),
         ("80% false accept  —  anchors on nearby text, cannot tell ",
          10.5, False, True, BODY, SANS)],
        [("   ",   10, False, False, DARK, SANS),
         ("\"test AUC 0.8412\" from \"train AUC 0.8934\" in prose.",
          10.5, False, True, BODY, SANS)],
    ],
    [[
        ("none produces an auditable scalar", 11, True, True, RED, SANS),
    ]],
)

# arrow 1
arrow(slide, l1 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), BLUE)

# ---------- Panel 2: The Geometric Judge (solution) ----------------------
l2 = l1 + panel_w + gap
make_panel(
    l2, BLUE, BLUE_BG,
    "Geometric Judge",
    "inverse provenance",
    [
        [("Every question was generated by ",
          11.5, False, False, BODY, SANS),
         ("one primitive", 11.5, True, False, DARK, SANS),
         (".  Score it with the ",
          11.5, False, False, BODY, SANS),
         ("same primitive", 11.5, True, False, DARK, SANS),
         (".",
          11.5, False, False, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Five measures, never averaged:",
          11.5, True, False, BLUE, SANS)],
        [("    ·  distance",  11, True, False, DARK, SANS),
         ("   plausibility, link, multi-hop, ...",
          10.5, False, True, BODY, SANS)],
        [("    ·  polarity",  11, True, False, DARK, SANS),
         ("   agree / unrelated / contradict",
          10.5, False, True, BODY, SANS)],
        [("    ·  path",      11, True, False, DARK, SANS),
         ("   multi-hop consistency",
          10.5, False, True, BODY, SANS)],
        [("    ·  phase",     11, True, False, DARK, SANS),
         ("   ordered numeric comparison",
          10.5, False, True, BODY, SANS)],
        [("    ·  exact",     11, True, False, DARK, SANS),
         ("   bit-exact value + integrity hash",
          10.5, False, True, BODY, SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Thresholds ",
          11, False, False, BODY, SANS),
         ("calibrated from the store itself",
          11, True, False, DARK, SANS),
         ("  —  no hand-tuning, no LLM help.",
          11, False, False, BODY, SANS)],
    ],
    [[
        ("one question  →  one primitive  →  one scalar",
         11, True, True, BLUE, SANS),
    ]],
)

# arrow 2
arrow(slide, l2 + panel_w + Inches(0.05), t_top + Inches(1.67),
      Inches(0.35), Inches(0.40), GREEN)

# ---------- Panel 3: Head-to-head (result) -------------------------------
l3 = l2 + panel_w + gap
make_panel(
    l3, GREEN, GREEN_BG,
    "Head-to-head",
    "same benchmark, same day",
    [
        [("           ", 10.5, False, False, DARK, MONO),
         ("Geometric",  10.5, True, False, GREEN, SANS),
         ("      ", 10.5, False, False, DARK, SANS),
         ("LLM-as-judge", 10.5, True, False, RED,   SANS)],
        [(" ", 4, False, False, DARK, SANS)],
        [("Accuracy",          11,   True,  False, DARK, SANS),
         ("         ",         11,   False, False, DARK, MONO),
         ("92.7%",              11.5, True,  False, GREEN, MONO),
         ("        ",           11,   False, False, DARK, MONO),
         ("53.1%",              11.5, True,  False, RED, MONO)],
        [("False accept",      11,   True,  False, DARK, SANS),
         ("     ",              11,   False, False, DARK, MONO),
         (" 7.7%",              11.5, True,  False, GREEN, MONO),
         ("        ",           11,   False, False, DARK, MONO),
         ("80.0%",              11.5, True,  False, RED, MONO)],
        [("Latency",           11,   True,  False, DARK, SANS),
         ("           ",       11,   False, False, DARK, MONO),
         (" < 1 s",             11.5, True,  False, GREEN, MONO),
         ("        ",           11,   False, False, DARK, MONO),
         ("~ 9 min",            11.5, True,  False, RED, MONO)],
        [(" ", 6, False, False, DARK, SANS)],
        [("Every verdict carries:",
          11, True, False, GREEN, SANS)],
        [("  ·  binary outcome   (accept / reject)",
          10.5, False, False, BODY, SANS)],
        [("  ·  continuous scalar  (for drift / margin)",
          10.5, False, False, BODY, SANS)],
        [("  ·  one-line audit detail  (primitive · score · threshold)",
          10.5, False, False, BODY, SANS)],
    ],
    [[
        ("deterministic  ·  auditable  ·  fast",
         11, True, True, GREEN, SANS),
    ]],
)

# -- Decision-detail example (small strip) --------------------------------
dd_top = Inches(5.70)
box(slide, Inches(0.5), dd_top, Inches(12.4), Inches(0.55), WHITE,
    line=FAINT, rounded=True)
txt(slide, Inches(0.7), dd_top + Inches(0.07), Inches(2.7), Inches(0.40), [[
    ("Example decision detail:", 11, True, False, NAVY, SANS),
]], anchor=MSO_ANCHOR.MIDDLE)
txt(slide, Inches(3.2), dd_top + Inches(0.07), Inches(9.5), Inches(0.40), [[
    ("[distance] plausibility   d = 1.056,  τ = 1.388,  P(real) = 0.99   →   ",
     11, False, False, DARK, MONO),
    ("ACCEPT",
     11, True, False, GREEN, MONO),
]], anchor=MSO_ANCHOR.MIDDLE)

# -- Bottom punchline bar -------------------------------------------------
bar_top = Inches(6.42)
box(slide, Inches(0.5), bar_top, Inches(12.4), Inches(0.95), NAVY, rounded=True)
txt(slide, Inches(0.75), bar_top + Inches(0.10), Inches(11.9), Inches(0.35), [[
    ("No LLM judgment enters scoring.  ",
     12, True, False, WHITE, SANS),
    ("Same seed  →  same verdicts.  Same verdicts  →  meaningful regression testing.",
     12, False, False, WHITE, SANS),
]])
txt(slide, Inches(0.75), bar_top + Inches(0.48), Inches(11.9), Inches(0.42), [[
    ("Each verdict is a calibrated scalar against a store-fit threshold  —  ",
     11, False, True, WHITE, SANS),
    ("a reviewer can reconstruct any decision from the one-line detail, without re-running the judge.",
     11, False, True, WHITE, SANS),
]])

out = "/home/asudjianto/jupyterlab/kg-memory/judge_slide.pptx"
prs.save(out)
print(f"Wrote: {out}")
