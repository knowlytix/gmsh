"""1-page slide: Building Ground Truth — From Text to Structured Knowledge.

Concept first: text -> structured graph -> verifiable ground truth.
Implementation details (parse/index/train/store) are labels, not paragraphs.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Colors --------------------------------------------------------------
NAVY   = RGBColor(0x14, 0x2B, 0x4D)
DARK   = RGBColor(0x22, 0x26, 0x2E)
BODY   = RGBColor(0x3E, 0x44, 0x50)
MID    = RGBColor(0x6C, 0x72, 0x80)
FAINT  = RGBColor(0xD3, 0xD7, 0xDE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

STEP1  = RGBColor(0x4A, 0x69, 0x8F)   # doc (cool slate)
STEP2  = RGBColor(0x1F, 0x6F, 0xB5)   # graph (primary blue)
STEP3  = RGBColor(0x2E, 0x8B, 0x57)   # ground truth (trust-green)

STEP1_BG = RGBColor(0xEC, 0xEF, 0xF5)
STEP2_BG = RGBColor(0xE6, 0xF0, 0xFA)
STEP3_BG = RGBColor(0xE7, 0xF4, 0xEC)

WARM   = RGBColor(0xC9, 0x56, 0x2D)
MONO   = "Consolas"
SANS   = "Calibri"


def txt(slide, left, top, width, height, runs,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        margin=(0.02, 0.02, 0.02, 0.02)):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left   = Inches(margin[0])
    tf.margin_right  = Inches(margin[1])
    tf.margin_top    = Inches(margin[2])
    tf.margin_bottom = Inches(margin[3])
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, italic, color, font in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = font
    return tb


def box(slide, left, top, width, height, fill, line=None, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        left, top, width, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def arrow(slide, left, top, width, height, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


# --- Slide ---------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Navy accent at top
box(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.10), NAVY)

# Title
txt(slide, Inches(0.5), Inches(0.22), Inches(12.4), Inches(0.55), [[
    ("Building Ground Truth", 30, True, False, NAVY, SANS),
    ("   —   from text to structured knowledge", 18, False, False, MID, SANS),
]])
txt(slide, Inches(0.5), Inches(0.78), Inches(12.4), Inches(0.28), [[
    ("Essay 2  ·  Agus Sudjianto  &  Aldo Lau", 11, False, True, MID, SANS),
]])

# One-line framing
txt(slide, Inches(0.5), Inches(1.18), Inches(12.4), Inches(0.38), [[
    ("Agents work on documents, not labeled datasets.  ",
     14, True, False, DARK, SANS),
    ("The ground truth must be the document, made verifiable.",
     14, False, False, BODY, SANS),
]])

# -- The core transformation (three big panels) ---------------------------
t_top   = Inches(1.85)
t_h     = Inches(3.30)
panel_w = Inches(3.85)
gap     = Inches(0.45)
total_w = 3 * panel_w + 2 * gap
start_l = Inches(0.5 + (12.4 - (3 * 3.85 + 2 * 0.45)) / 2.0)

# Panel 1 — Document (prose)
l1 = start_l
box(slide, l1, t_top, panel_w, t_h, fill=STEP1_BG, line=STEP1, rounded=True)
box(slide, l1, t_top, panel_w, Inches(0.45), fill=STEP1, rounded=True)
box(slide, l1, t_top + Inches(0.33), panel_w, Inches(0.12), fill=STEP1)
txt(slide, l1 + Inches(0.15), t_top + Inches(0.04),
    panel_w - Inches(0.3), Inches(0.38), [[
        ("Document", 16, True, False, WHITE, SANS),
        ("   prose, tables, bullets", 11, False, True, WHITE, SANS),
    ]], anchor=MSO_ANCHOR.MIDDLE)
# sample text
box(slide, l1 + Inches(0.25), t_top + Inches(0.70),
    panel_w - Inches(0.5), Inches(2.1), fill=WHITE, line=FAINT, rounded=True)
txt(slide, l1 + Inches(0.40), t_top + Inches(0.85),
    panel_w - Inches(0.80), Inches(1.85),
    [
        [("“...  MDL-001  achieved  ", 11, False, False, BODY, MONO)],
        [("train AUC 0.8934, ",  11, True,  False, DARK, MONO)],
        [("test AUC 0.8412, ",   11, True,  False, DARK, MONO)],
        [("with a train-test gap ", 11, False, False, BODY, MONO)],
        [("of 0.0522.  ",        11, True,  False, DARK, MONO)],
        [("Depends on MDL-009. ”", 11, False, False, BODY, MONO)],
    ])
txt(slide, l1 + Inches(0.20), t_top + t_h - Inches(0.42),
    panel_w - Inches(0.4), Inches(0.35), [[
        ("ambiguous  ·  not queryable  ·  prose",
         11, False, True, STEP1, SANS),
    ]], align=PP_ALIGN.CENTER)

# Arrow 1
arrow(slide, l1 + panel_w + Inches(0.05), t_top + Inches(1.45),
      Inches(0.35), Inches(0.40), STEP2)

# Panel 2 — Knowledge graph
l2 = l1 + panel_w + gap
box(slide, l2, t_top, panel_w, t_h, fill=STEP2_BG, line=STEP2, rounded=True)
box(slide, l2, t_top, panel_w, Inches(0.45), fill=STEP2, rounded=True)
box(slide, l2, t_top + Inches(0.33), panel_w, Inches(0.12), fill=STEP2)
txt(slide, l2 + Inches(0.15), t_top + Inches(0.04),
    panel_w - Inches(0.3), Inches(0.38), [[
        ("Knowledge Graph", 16, True, False, WHITE, SANS),
        ("   head · relation · tail", 11, False, True, WHITE, SANS),
    ]], anchor=MSO_ANCHOR.MIDDLE)
box(slide, l2 + Inches(0.25), t_top + Inches(0.70),
    panel_w - Inches(0.5), Inches(2.1), fill=WHITE, line=FAINT, rounded=True)
txt(slide, l2 + Inches(0.35), t_top + Inches(0.82),
    panel_w - Inches(0.7), Inches(1.95),
    [
        [("(MDL-001, has_train_auc, 0.8934)", 10.5, False, False, DARK, MONO)],
        [("(MDL-001, has_test_auc,  0.8412)", 10.5, True,  False, STEP2, MONO)],
        [("(MDL-001, has_gap,       0.0522)", 10.5, False, False, DARK, MONO)],
        [("(MDL-001, depends_on,    MDL-009)", 10.5, False, False, DARK, MONO)],
        [("...", 10.5, False, True, MID, MONO)],
    ])
txt(slide, l2 + Inches(0.20), t_top + t_h - Inches(0.42),
    panel_w - Inches(0.4), Inches(0.35), [[
        ("explicit  ·  typed  ·  exact-match queryable",
         11, False, True, STEP2, SANS),
    ]], align=PP_ALIGN.CENTER)

# Arrow 2
arrow(slide, l2 + panel_w + Inches(0.05), t_top + Inches(1.45),
      Inches(0.35), Inches(0.40), STEP3)

# Panel 3 — Ground-truth store
l3 = l2 + panel_w + gap
box(slide, l3, t_top, panel_w, t_h, fill=STEP3_BG, line=STEP3, rounded=True)
box(slide, l3, t_top, panel_w, Inches(0.45), fill=STEP3, rounded=True)
box(slide, l3, t_top + Inches(0.33), panel_w, Inches(0.12), fill=STEP3)
txt(slide, l3 + Inches(0.15), t_top + Inches(0.04),
    panel_w - Inches(0.3), Inches(0.38), [[
        ("Ground-Truth Store", 16, True, False, WHITE, SANS),
        ("   verifiable", 11, False, True, WHITE, SANS),
    ]], anchor=MSO_ANCHOR.MIDDLE)
box(slide, l3 + Inches(0.25), t_top + Inches(0.70),
    panel_w - Inches(0.5), Inches(2.1), fill=WHITE, line=FAINT, rounded=True)
txt(slide, l3 + Inches(0.40), t_top + Inches(0.85),
    panel_w - Inches(0.8), Inches(1.95),
    [
        [("Geometric embedding",  12, True, False, DARK, SANS)],
        [("    plausibility = distance",  11, False, False, BODY, SANS)],
        [("    contradiction = a scalar", 11, False, False, BODY, SANS)],
        [(" ",                          6,  False, False, DARK, SANS)],
        [("Exact numeric register",     12, True, False, DARK, SANS)],
        [("    bit-exact values + integrity hash",
          11, False, False, BODY, SANS)],
    ])
txt(slide, l3 + Inches(0.20), t_top + t_h - Inches(0.42),
    panel_w - Inches(0.4), Inches(0.35), [[
        ("every claim  →  one deterministic check",
         11, False, True, STEP3, SANS),
    ]], align=PP_ALIGN.CENTER)

# -- Four steps underneath (labels, not paragraphs) -----------------------
steps_top = Inches(5.35)
txt(slide, Inches(0.5), steps_top, Inches(12.4), Inches(0.28), [[
    ("Four stages of the pipeline:",
     12, True, False, MID, SANS),
    ("    Parse  →  Index  →  Train  →  Store",
     13, True, False, NAVY, MONO),
]])
txt(slide, Inches(0.5), steps_top + Inches(0.32), Inches(12.4), Inches(0.30), [[
    ("read tables & bullets  ·  build typed vocabulary  ·  learn geometry  ·  hash exact numbers",
     11, False, True, MID, SANS),
]])

# -- What it unlocks (condensed, single row) ------------------------------
unlock_top = Inches(6.10)
txt(slide, Inches(0.5), unlock_top, Inches(12.4), Inches(0.30), [[
    ("Every agent claim becomes a deterministic check:",
     12, True, False, WARM, SANS),
]])
unlock = [
    "fact supported?  distance",
    "contradicts?  scalar",
    "reasoning consistent?  path check",
    "exact number?  lookup + hash",
    "meets threshold?  encoded compare",
]
u_row_top = Inches(6.42)
slot_w = Inches(12.4 / 5.0)
for i, item in enumerate(unlock):
    left = Inches(0.5) + slot_w * i
    # tiny warm bullet
    box(slide, left + Inches(0.05), u_row_top + Inches(0.10),
        Inches(0.08), Inches(0.08), WARM, rounded=True)
    txt(slide, left + Inches(0.20), u_row_top, slot_w - Inches(0.25), Inches(0.30),
        [[(item, 10.5, False, False, DARK, SANS)]])

# -- Bottom result bar ----------------------------------------------------
bar_top = Inches(6.85)
box(slide, Inches(0.5), bar_top, Inches(12.4), Inches(0.50), NAVY, rounded=True)
txt(slide, Inches(0.75), bar_top + Inches(0.10), Inches(11.9), Inches(0.32), [[
    ("14,544-char risk report  →  ", 11.5, True, False, WHITE, SANS),
    ("508 entities  ·  66 relations  ·  807 triples  ·  261 exact numbers.",
     11.5, False, False, WHITE, SANS),
    ("   No prose left to re-interpret.",
     11.5, False, True, WHITE, SANS),
]])

out = "/home/asudjianto/jupyterlab/kg-memory/ground_truth_slide.pptx"
prs.save(out)
print(f"Wrote: {out}")
